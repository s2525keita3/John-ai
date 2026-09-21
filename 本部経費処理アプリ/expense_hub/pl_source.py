"""
PL（収支計画スプシ）の読み口。xlsx でもスプシ直結でも同じ Grid に揃える。
Grid.cell(row, col) -> (値, メモ)。行・列は1始まり（openpyxl と同じ）。

スプシ直結は Sheets API の読み取り専用スコープだけを使う。PLへは書き込まない。
認証＝Streamlit secrets の [gcp_service_account]（サービスアカウントJSONの中身）。
対象スプシはそのサービスアカウントのメールに「閲覧者」で共有しておく。
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

import openpyxl
from openpyxl.utils import get_column_letter

PL_SPREADSHEET_ID = "1rPs01xlB1Iv8a8ovRH8eSIDJGvqCa1mn5SLM2SwL71A"  # 2026年 ステーション収支計画
SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]


@dataclass
class Grid:
    title: str
    cells: dict[tuple[int, int], tuple[object, str]] = field(default_factory=dict)
    max_row: int = 0
    max_column: int = 0
    origin: str = ""

    def cell(self, r: int, c: int) -> tuple[object, str]:
        return self.cells.get((r, c), (None, ""))

    def value(self, r: int, c: int):
        return self.cell(r, c)[0]

    def note(self, r: int, c: int) -> str:
        return self.cell(r, c)[1]

    @staticmethod
    def coord(r: int, c: int) -> str:
        return f"{get_column_letter(c)}{r}"

    def to_xlsx(self) -> bytes:
        """バックアップ用：値（数式はそのまま）とメモを1枚のxlsxに書き出す。"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = self.title[:31]
        for (r, c), (v, note) in self.cells.items():
            cell = ws.cell(r, c, v)
            if note:
                cell.comment = openpyxl.comments.Comment(note, "backup")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def resolve_titles(available: list[str], wanted: list[str]) -> dict[str, str]:
    """欲しいタブ名 → 実物のタブ名。前後の空白違い（例「 桜：収支」）を吸収する。無いものは入れない。"""
    out = {}
    for w in wanted:
        hit = next((t for t in available if t == w), None) or next((t for t in available if t.strip() == w.strip()), None)
        if hit:
            out[w] = hit
    return out


def _grid_from_ws(ws, title: str) -> Grid:
    g = Grid(title, max_row=ws.max_row, max_column=ws.max_column, origin="xlsx")
    for row in ws.iter_rows():
        for c in row:
            note = c.comment.text if c.comment else ""
            if c.value is not None or note:
                g.cells[(c.row, c.column)] = (c.value, note)
    return g


def grids_from_xlsx(xlsx_bytes: bytes, sheets: list[str]) -> dict[str, Grid]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    names = resolve_titles(wb.sheetnames, sheets)
    return {w: _grid_from_ws(wb[t], w) for w, t in names.items()}


def grid_from_xlsx(xlsx_bytes: bytes, sheet: str) -> Grid:
    g = grids_from_xlsx(xlsx_bytes, [sheet]).get(sheet)
    if g is None:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
        raise KeyError(f"タブ「{sheet}」がありません（あるタブ：{', '.join(wb.sheetnames)}）")
    return g


def _sheets_service(sa_info: dict):
    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build

    creds = Credentials.from_service_account_info(sa_info, scopes=SCOPES)
    return build("sheets", "v4", credentials=creds, cache_discovery=False)


def list_sheet_titles(sa_info: dict, spreadsheet_id: str = PL_SPREADSHEET_ID) -> list[str]:
    meta = _sheets_service(sa_info).spreadsheets().get(
        spreadsheetId=spreadsheet_id, fields="sheets.properties.title"
    ).execute()
    return [s["properties"]["title"] for s in meta.get("sheets", [])]


def grids_from_sheets(sa_info: dict, sheets: list[str], spreadsheet_id: str = PL_SPREADSHEET_ID) -> dict[str, Grid]:
    """複数タブを1回で取る。セルの入力値（数式はそのまま）・計算後の値・メモ。"""
    names = resolve_titles(list_sheet_titles(sa_info, spreadsheet_id), sheets)
    if not names:
        return {}
    back = {t: w for w, t in names.items()}
    resp = _sheets_service(sa_info).spreadsheets().get(
        spreadsheetId=spreadsheet_id,
        ranges=[f"'{t}'" for t in names.values()],
        includeGridData=True,
        fields="sheets(properties(title),data(startRow,startColumn,rowData(values(note,userEnteredValue,effectiveValue))))",
    ).execute()
    out = {}
    for sh in resp.get("sheets", []):
        title = sh["properties"]["title"]
        g = Grid(back.get(title, title), origin="sheets")
        for block in sh.get("data", []):
            r0 = block.get("startRow", 0)
            c0 = block.get("startColumn", 0)
            for i, row in enumerate(block.get("rowData", []) or []):
                for j, v in enumerate(row.get("values", []) or []):
                    note = v.get("note", "")
                    uev = v.get("userEnteredValue", {})
                    ev = v.get("effectiveValue", {})
                    if "formulaValue" in uev:
                        val = uev["formulaValue"]
                    elif "numberValue" in ev:
                        val = ev["numberValue"]
                    elif "stringValue" in ev:
                        val = ev["stringValue"]
                    elif "boolValue" in ev:
                        val = ev["boolValue"]
                    else:
                        val = None
                    if val is None and not note:
                        continue
                    r, c = r0 + i + 1, c0 + j + 1
                    g.cells[(r, c)] = (val, note)
                    g.max_row = max(g.max_row, r)
                    g.max_column = max(g.max_column, c)
        out[g.title] = g
    return out


def grid_from_sheets(sa_info: dict, sheet: str, spreadsheet_id: str = PL_SPREADSHEET_ID) -> Grid:
    g = grids_from_sheets(sa_info, [sheet], spreadsheet_id).get(sheet)
    if g is None:
        raise KeyError(f"タブ「{sheet}」が取れませんでした")
    return g
