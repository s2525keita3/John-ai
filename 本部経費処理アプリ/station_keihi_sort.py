"""
店舗（ステーション）経費の自動仕分け — 桜木町から適用。

入力:
  1) あおぞらネット銀行 標準CSV（日付・摘要・入金金額・出金金額・残高・メモ）
     または 経費Excel「YY.MM」シートの銀行表（日付・科目・支払先・摘要・入金・出金・計）
  2) 小口（Googleスプシ「桜木町　小口」→xlsx/csv。日付・科目・スタッフ・支払先・摘要・入金・出金・計）
  3) キーワードマスタ（sakuragicho_master.csv）

出力:
  - 行ごとの 店舗科目 / PL行 / 判定（確定・要確認・判断不能・要按分）/ 根拠
  - 経費Excel末尾と同じ「科目×銀行/小口/計」サマリ
  - 収支計画スプシ「 桜：収支」に貼れる PL行別の実績

依存: pandas（openpyxl は xlsx 読み込み時のみ）
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

from station_accounts import (
    PL_NEEDS_SPLIT,
    PL_OUT_OF_SCOPE,
    PL_ROWS,
    STATION_ACCOUNTS,
    STATION_TO_PL,
    WELFARE_SPLIT_RULES,
    normalize_account,
)

HERE = Path(__file__).resolve().parent
DEFAULT_MASTER = HERE / "sakuragicho_master.csv"

COLS = ("日付", "科目", "スタッフ", "支払先", "摘要", "入金", "出金")


# ---------------------------------------------------------------------------
# 正規化（本部経費処理アプリ classifier.normalize_for_match と同じ規則）
# ---------------------------------------------------------------------------
def norm(s) -> str:
    s = unicodedata.normalize("NFKC", str(s or ""))
    s = re.sub(r"\s+", "", s)
    return s.upper()


def to_num(v) -> float:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("¥", "").replace("円", "")
    if s in ("", "-", "—"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


# ---------------------------------------------------------------------------
# マスタ
# ---------------------------------------------------------------------------
@dataclass
class Rule:
    keyword: str
    account: str
    scope: str        # 銀行 / 小口 / 共通
    status: str       # 実績確認済 / カナ表記要確認 / 要確認
    note: str

    @property
    def key(self) -> str:
        return norm(self.keyword)


def load_master(path: str | Path = DEFAULT_MASTER) -> list[Rule]:
    df = pd.read_csv(path, encoding="utf-8-sig").fillna("")
    rules = [
        Rule(
            keyword=str(r["摘要キーワード"]).strip(),
            account=str(r["店舗科目"]).strip(),
            scope=str(r.get("区分", "共通")).strip() or "共通",
            status=str(r.get("確認状況", "")).strip(),
            note=str(r.get("備考", "")).strip(),
        )
        for _, r in df.iterrows()
        if str(r["摘要キーワード"]).strip() and str(r["店舗科目"]).strip()
    ]
    # 長いキーワード優先（「SMBC賃料」＞「SMBC」、「代表口座」の長文が最優先）
    rules.sort(key=lambda r: len(r.key), reverse=True)
    return rules


# ---------------------------------------------------------------------------
# 入力の読み込み（列名ゆれ吸収）
# ---------------------------------------------------------------------------
_COLMAP = {
    "日付": ("日付", "取引日", "date"),
    "科目": ("科目", "勘定科目", "振分PL項目", "PL"),
    "スタッフ": ("スタッフ", "担当", "氏名", "名字"),
    "支払先": ("支払先", "取引先", "相手先"),
    "摘要": ("摘要", "内容", "備考", "メモ", "ご利用内容"),
    "入金": ("入金", "入金額", "入金金額"),
    "出金": ("出金", "出金額", "出金金額", "支出", "支払"),
}


def standardize(df: pd.DataFrame) -> pd.DataFrame:
    """任意の明細表を COLS へ寄せる。あおぞらCSVは 支払先 が無く 摘要 だけ。"""
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    ren: dict[str, str] = {}
    for std, cands in _COLMAP.items():
        for c in df.columns:
            if c in ren:
                continue
            if c == std or c in cands:
                ren[c] = std
                break
    df = df.rename(columns=ren)
    for c in COLS:
        if c not in df.columns:
            df[c] = ""
    df["入金"] = df["入金"].map(to_num)
    df["出金"] = df["出金"].map(to_num)
    df["科目"] = df["科目"].map(normalize_account)
    for c in ("スタッフ", "支払先", "摘要"):
        df[c] = df[c].fillna("").astype(str).str.strip()
    df = df[(df["入金"] != 0) | (df["出金"] != 0)]
    df = df[(df["支払先"] != "") | (df["摘要"] != "")]
    return df[list(COLS)].reset_index(drop=True)


def read_aozora_csv(path_or_buf) -> pd.DataFrame:
    """GMOあおぞら 標準CSV（複数エンコーディングを試す）。"""
    last = None
    for enc in ("utf-8-sig", "cp932", "utf-8", "utf-16"):
        try:
            return standardize(pd.read_csv(path_or_buf, encoding=enc))
        except (UnicodeDecodeError, UnicodeError) as e:  # pragma: no cover
            last = e
            if hasattr(path_or_buf, "seek"):
                path_or_buf.seek(0)
    raise last or ValueError("CSVを読めません")


def read_keihi_excel_bank(path, sheet: str) -> pd.DataFrame:
    """経費Excel『YY.MM』シートの銀行表（A〜G列）。"""
    raw = pd.read_excel(path, sheet_name=sheet, header=None, usecols="A:G", skiprows=3)
    raw.columns = list(COLS[:2]) + ["支払先", "摘要", "入金", "出金", "計"]
    raw["スタッフ"] = ""
    return standardize(raw)


def read_petty(path, sheet: str | None = None) -> pd.DataFrame:
    """小口（xlsx/csv）。ヘッダ行は「日付」を含む行を自動検出。"""
    p = str(path)
    if p.lower().endswith((".xlsx", ".xlsm")):
        raw = pd.read_excel(path, sheet_name=sheet or 0, header=None)
    else:
        raw = pd.read_csv(path, header=None, encoding="utf-8-sig")
    hdr = None
    for i in range(min(10, len(raw))):
        vals = [str(v).strip() for v in raw.iloc[i].tolist()]
        if "日付" in vals and ("科目" in vals or "摘要" in vals):
            hdr = i
            break
    if hdr is None:
        raise ValueError("小口: ヘッダ行（日付・科目…）が見つかりません")
    df = raw.iloc[hdr + 1 :].copy()
    df.columns = [str(v).strip() for v in raw.iloc[hdr].tolist()]
    df = df.loc[:, [c for c in df.columns if c and c != "nan"]]
    return standardize(df)


# ---------------------------------------------------------------------------
# 仕分け本体
# ---------------------------------------------------------------------------
_USER_SUFFIX = re.compile(r"(様|サマ)$")
_GENERIC_SPLIT_HINTS: tuple[tuple[str, str], ...] = (
    # 代表口座（本部相殺）一括振込の中身。摘要にこの語があれば店舗科目まで確定できる。
    # 上から順に評価（「本部管理費から相殺」の注記に負けないよう、駐車場・点検を先に置く）
    ("駐車場", "賃借料"),
    ("点検", "車両費"),
    ("車検", "車両費"),
    ("所得税", "福利厚生費"),
    ("市民税", "法人税"),
    ("住民税", "法人税"),
    ("社会保険料", "福利厚生費"),
    ("中小企業退職金共済", "福利厚生費"),
    ("中退共", "福利厚生費"),
    ("SOFTBANK", "通信費"),
    ("ソフトバンク", "通信費"),
    ("朝日ネット", "通信費"),
    ("LINEWORKS", "通信費"),
    ("日新火災", "保険料"),
    ("エネクスフリート", "旅費交通費"),
    ("燃料", "旅費交通費"),
    ("飲み会", "福利厚生費"),
    ("懇親", "福利厚生費"),
    ("本部管理費", "本部経費"),
)

# 通販（Amazon等）の品名ヒント：スタッフ用の飲料・菓子は福利厚生費（経費Excel 26.07 実績）
_WELFARE_ITEM_HINTS = ("天然水", "水500", "水2L", "麦茶", "緑茶", "お茶", "コーヒー", "飲料", "ジュース",
                       "菓子", "スイーツ", "ケーキ", "弁当", "飲み物")


def _pl_for(account: str, text_norm: str) -> str:
    if account == "要按分":
        return PL_NEEDS_SPLIT
    if account == "福利厚生費":
        for kw, pl in WELFARE_SPLIT_RULES:
            if norm(kw) in text_norm:
                return pl
    return STATION_TO_PL.get(account, "")


def classify_row(
    row: pd.Series, rules: list[Rule], source: str
) -> tuple[str, str, str, str]:
    """
    返り値: (店舗科目, PL行, 判定, 根拠)
    source: '銀行' or '小口'
    判定: 確定 / 要確認 / 判断不能 / 要按分 / 入力済（小口等で科目が既に入っている）
    """
    payee = str(row.get("支払先", "") or "")
    memo = str(row.get("摘要", "") or "")
    text = f"{payee} {memo}"
    tn = norm(text)
    pn = norm(payee)
    is_in = to_num(row.get("入金")) > 0 and to_num(row.get("出金")) == 0
    pre = normalize_account(row.get("科目"))

    # 0) 入力済みの科目（小口シートなど）は尊重し、マスタと不一致なら要確認で知らせる
    if pre and pre in STATION_ACCOUNTS:
        pl = _pl_for(pre, tn)
        hit = _first_hit(tn, rules, source, pn)
        if hit and hit.account not in (pre, "要按分") and not (is_in and pre == "入金"):
            return pre, pl, "要確認", f"入力済『{pre}』だがマスタは『{hit.account}』（{hit.keyword}）"
        if pre not in ("入金",) and to_num(row.get("入金")) > 0 and "相殺" not in tn and "返金" not in tn:
            return pre, pl, "要確認", f"経費科目『{pre}』の行に入金額がある（補充・立替精算の入力ミス？）"
        return pre, pl, "入力済", "科目列の値を採用"

    # 1) 本部 代表口座との相殺一括振込 → 摘要に中身の語があれば確定、無ければ要按分
    hit = _first_hit(tn, rules, source, pn)
    if hit and hit.account == "要按分":
        for kw, acc in _GENERIC_SPLIT_HINTS:
            if norm(kw) in tn:
                pl = _pl_for(acc, tn)
                return acc, pl, "確定", f"本部相殺明細『{kw}』→{acc}"
        return "要按分", PL_NEEDS_SPLIT, "要按分", "本部相殺の一括振込。本部管理費明細で分解が必要"

    # 2) 利用者個人からの入金（「○○様」）— 名前はマスタに載せない（個人情報）
    if is_in and (_USER_SUFFIX.search(payee.strip()) or "利用料" in tn or "負担金" in tn):
        return "入金", "入金", "確定", "利用者負担金（個人名＋様／利用料）"

    # 3) マスタ一致
    if hit:
        acc = hit.account
        # 経費先からの入金（取消・返金・相殺）は入金へ
        if is_in and acc not in ("入金",):
            if acc in ("給料賃金",):
                return "入金", "入金", "要確認", f"給与名義からの入金（{hit.keyword}）— 返金？"
            return "入金", "入金", "確定", f"{hit.keyword}（{acc}）からの入金＝返金・相殺"
        # GMOあおぞら名義の出金は資金移動
        if (not is_in) and acc == "入金" and "GMO" in pn:
            return "出金", PL_OUT_OF_SCOPE, "確定", "自社口座間の資金移動"
        if acc == "備品・消耗品費" and "AMAZON" in pn and any(norm(k) in tn for k in _WELFARE_ITEM_HINTS):
            return "福利厚生費", "福利厚生", "確定", f"{hit.keyword}＋品名が飲料・菓子＝スタッフ福利厚生"
        pl = _pl_for(acc, tn)
        verdict = "確定" if hit.status == "実績確認済" else "要確認"
        why = f"キーワード『{hit.keyword}』"
        if hit.status != "実績確認済":
            why += f"（{hit.status}）"
        if hit.note and hit.status == "要確認":
            why += f" {hit.note}"
        return acc, pl, verdict, why

    # 4) 一致なし
    if is_in:
        return "入金", "入金", "要確認", "入金だがマスタ未登録（利用者個人・保険金・返金の可能性）"
    return "", "", "判断不能", "マスタに一致するキーワードなし"


def _first_hit(tn: str, rules: list[Rule], source: str, pn: str | None = None) -> Rule | None:
    """長いキーワード優先で最初に当たったルール。給与名義（給料賃金）は支払先だけで判定する
    （摘要の「名刺(○○)」「○○ 放置違反金」等でスタッフ名が出ても給与にしない）。"""
    for r in rules:
        if r.scope != "共通" and r.scope != source:
            continue
        if not r.key:
            continue
        if r.account == "給料賃金":
            if pn is not None and r.key in pn:
                return r
            if pn is None and r.key in tn:
                return r
            continue
        if r.key in tn:
            return r
    return None


def classify(df: pd.DataFrame, rules: list[Rule], source: str) -> pd.DataFrame:
    out = standardize(df) if list(df.columns) != list(COLS) else df.copy()
    res = [classify_row(r, rules, source) for _, r in out.iterrows()]
    out["店舗科目"] = [a for a, _, _, _ in res]
    out["PL行"] = [p for _, p, _, _ in res]
    out["判定"] = [v for _, _, v, _ in res]
    out["根拠"] = [w for _, _, _, w in res]
    out["ソース"] = source
    return out


# ---------------------------------------------------------------------------
# サマリ（経費Excel R〜U列と同じ形）と PL行別実績
# ---------------------------------------------------------------------------
def account_summary(bank: pd.DataFrame | None, petty: pd.DataFrame | None) -> pd.DataFrame:
    """科目×銀行/小口/計。入金は入金額、それ以外は 出金−入金（相殺入金を差し引く：経費Excel慣行）。"""
    order = [a for a in STATION_ACCOUNTS] + ["要按分", ""]
    rows = []
    for acc in order:
        rec = {"科目": acc if acc else "（判断不能）"}
        for name, d in (("銀行", bank), ("小口", petty)):
            if d is None or d.empty:
                rec[name] = 0.0
                continue
            sub = d[d["店舗科目"] == acc]
            if acc == "入金":
                rec[name] = float(sub["入金"].sum())
            else:
                rec[name] = float(sub["出金"].sum() - sub["入金"].sum())
        rec["計"] = rec["銀行"] + rec["小口"]
        if rec["計"] != 0 or acc in ("入金", "給料賃金", "本部経費", "旅費交通費", "賃借料", "通信費"):
            rows.append(rec)
    return pd.DataFrame(rows, columns=["科目", "銀行", "小口", "計"])


def pl_summary(bank: pd.DataFrame | None, petty: pd.DataFrame | None) -> pd.DataFrame:
    """収支計画スプシ「 桜：収支」の行順で実績を出す。人件費・賞与は支給控除一覧が正なので参考値。"""
    parts = [d for d in (bank, petty) if d is not None and not d.empty]
    if not parts:
        return pd.DataFrame(columns=["PL行", "金額", "備考"])
    all_ = pd.concat(parts, ignore_index=True)
    rows = []
    for pl in PL_ROWS:
        sub = all_[all_["PL行"] == pl]
        amt = float(sub["入金"].sum()) if pl == "入金" else float(sub["出金"].sum() - sub["入金"].sum())
        note = ""
        if pl.startswith("人件費") or pl == "賞与":
            note = "参考値（正は支給控除一覧）"
        if pl == "法人税":
            note = "店舗Excelの『法人税』は市民税＝預り金のためPL対象外。ここは0"
        rows.append({"PL行": pl, "金額": amt, "備考": note})
    split = all_[all_["PL行"] == PL_NEEDS_SPLIT]
    if not split.empty:
        rows.append({"PL行": PL_NEEDS_SPLIT, "金額": float(split["出金"].sum() - split["入金"].sum()),
                     "備考": f"{len(split)}行。本部管理費明細で分解してから上の行へ"})
    oos = all_[all_["PL行"] == PL_OUT_OF_SCOPE]
    if not oos.empty:
        rows.append({"PL行": PL_OUT_OF_SCOPE, "金額": float(oos["出金"].sum() - oos["入金"].sum()),
                     "備考": "小口補充・資金移動・預り金納付（所得税/市民税）"})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main(argv: list[str] | None = None) -> int:
    import argparse

    ap = argparse.ArgumentParser(description="店舗経費 自動仕分け（桜木町）")
    ap.add_argument("--bank-csv", help="あおぞら標準CSV")
    ap.add_argument("--keihi-xlsx", help="経費Excel（銀行表を読む場合）")
    ap.add_argument("--keihi-sheet", help="経費Excelのシート名 例: 26.08")
    ap.add_argument("--petty", help="小口 xlsx/csv")
    ap.add_argument("--petty-sheet", help="小口のシート名 例: 26.08")
    ap.add_argument("--master", default=str(DEFAULT_MASTER))
    ap.add_argument("--out", default="station_keihi_result.xlsx")
    a = ap.parse_args(argv)

    rules = load_master(a.master)
    bank = petty = None
    if a.bank_csv:
        bank = classify(read_aozora_csv(a.bank_csv), rules, "銀行")
    elif a.keihi_xlsx and a.keihi_sheet:
        bank = classify(read_keihi_excel_bank(a.keihi_xlsx, a.keihi_sheet), rules, "銀行")
    if a.petty:
        petty = classify(read_petty(a.petty, a.petty_sheet), rules, "小口")

    with pd.ExcelWriter(a.out) as xw:
        if bank is not None:
            bank.to_excel(xw, sheet_name="銀行_仕分け", index=False)
        if petty is not None:
            petty.to_excel(xw, sheet_name="小口_仕分け", index=False)
        account_summary(bank, petty).to_excel(xw, sheet_name="科目サマリ", index=False)
        pl_summary(bank, petty).to_excel(xw, sheet_name="PL行", index=False)
        pend = pd.concat([d[d["判定"].isin(("要確認", "判断不能", "要按分"))]
                          for d in (bank, petty) if d is not None], ignore_index=True)
        pend.to_excel(xw, sheet_name="要確認", index=False)
    print(f"書き出し: {a.out}")
    for name, d in (("銀行", bank), ("小口", petty)):
        if d is not None:
            print(name, d["判定"].value_counts().to_dict())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


# ===========================================================================
# 追加機能（2026-09-07 第2弾）
#  1) 本部管理費明細CSVで「要按分」一括振込を分解
#  2) 経費Excel「YY.MM」シート形式のワークブックを生成（A〜G 銀行／I〜P 小口／R〜U サマリ）
#  3) 収支計画スプシ「 桜：収支」への貼り付けブロック（実績列＋セルメモ文）
# ===========================================================================
import io as _io


def read_hq_detail(path_or_buf) -> pd.DataFrame:
    """本部管理費明細（本部→店舗の相殺一括振込の内訳）。列: 内容（or 摘要/項目）, 金額（or 出金/入金）。"""
    p = str(getattr(path_or_buf, "name", path_or_buf))
    if p.lower().endswith((".xlsx", ".xlsm")):
        df = pd.read_excel(path_or_buf)
    else:
        df = None
        for enc in ("utf-8-sig", "cp932", "utf-8"):
            try:
                if hasattr(path_or_buf, "seek"):
                    path_or_buf.seek(0)
                df = pd.read_csv(path_or_buf, encoding=enc)
                break
            except (UnicodeDecodeError, UnicodeError):
                continue
        if df is None:
            raise ValueError("本部管理費明細を読めません")
    df.columns = [str(c).strip() for c in df.columns]
    ren = {}
    for c in df.columns:
        if c in ("内容", "摘要", "項目", "明細") and "内容" not in ren.values():
            ren[c] = "内容"
        elif c in ("金額", "出金", "出金額", "請求額") and "金額" not in ren.values():
            ren[c] = "金額"
        elif c in ("入金", "入金額", "相殺", "控除") and "入金" not in ren.values():
            ren[c] = "入金"
        elif c in ("日付", "取引日"):
            ren[c] = "日付"
    df = df.rename(columns=ren)
    if "内容" not in df.columns or "金額" not in df.columns:
        raise ValueError(f"本部管理費明細: 「内容」「金額」列が必要です（現在: {list(df.columns)}）")
    df["金額"] = df["金額"].map(to_num)
    df["入金"] = df["入金"].map(to_num) if "入金" in df.columns else 0.0
    if "日付" not in df.columns:
        df["日付"] = ""
    df["内容"] = df["内容"].fillna("").astype(str).str.strip()
    return df[df["内容"] != ""][["日付", "内容", "金額", "入金"]].reset_index(drop=True)


def split_hq_lump(classified_bank: pd.DataFrame, detail: pd.DataFrame, rules: list[Rule]) -> tuple[pd.DataFrame, list[str]]:
    """
    分類済み銀行表の『要按分』行を、本部管理費明細で分解して差し替える。
    返り値: (差し替え後の銀行表, 警告メッセージ)
    - 明細の合計（金額−入金）と一括振込額が一致しなければ警告（分解はする）
    - 明細の各行は _GENERIC_SPLIT_HINTS → マスタ の順で科目を決める。決まらなければ要確認
    """
    warns: list[str] = []
    lump = classified_bank[classified_bank["判定"] == "要按分"]
    if lump.empty:
        return classified_bank, ["要按分の行はありません（明細は使いませんでした）"]
    if detail is None or detail.empty:
        return classified_bank, ["本部管理費明細が空です"]
    lump_total = float(lump["出金"].sum() - lump["入金"].sum())
    det_total = float(detail["金額"].sum() - detail["入金"].sum())
    if abs(lump_total - det_total) > 0.5:
        warns.append(f"一括振込 {lump_total:,.0f}円 と明細合計 {det_total:,.0f}円 が {lump_total - det_total:+,.0f}円 ずれています")
    base_date = str(lump.iloc[0]["日付"])
    payee = str(lump.iloc[0]["支払先"]) or "株式会社ジョン　GMOあおネット銀行　代表口座"
    new_rows = []
    for _, d in detail.iterrows():
        text = d["内容"]
        tn = norm(text)
        acc = ""
        for kw, a in _GENERIC_SPLIT_HINTS:
            if norm(kw) in tn:
                acc = a
                break
        if not acc:
            hit = _first_hit(tn, rules, "銀行")
            if hit and hit.account not in ("要按分", "給料賃金"):
                acc = hit.account
        is_in = float(d["入金"]) > 0 and float(d["金額"]) == 0
        row = {
            "日付": d["日付"] or base_date, "科目": acc, "スタッフ": "", "支払先": payee, "摘要": text,
            "入金": float(d["入金"]), "出金": float(d["金額"]),
        }
        if acc:
            pl = _pl_for(acc, tn)
            row.update({"店舗科目": acc, "PL行": pl, "判定": "確定", "根拠": "本部管理費明細で分解"})
        else:
            row.update({"店舗科目": "", "PL行": "", "判定": "要確認", "根拠": "本部管理費明細の科目が決められない"})
        row["ソース"] = "銀行"
        new_rows.append(row)
    keep = classified_bank[classified_bank["判定"] != "要按分"]
    out = pd.concat([keep, pd.DataFrame(new_rows)], ignore_index=True)
    if "日付" in out.columns:
        out = out.assign(_d=out["日付"].astype(str)).sort_values("_d", kind="stable").drop(columns="_d").reset_index(drop=True)
    warns.append(f"要按分 {len(lump)}行 → 明細 {len(new_rows)}行に分解（要確認 {sum(1 for r in new_rows if r['判定'] == '要確認')}行）")
    return out, warns


# ---- 経費Excel 月次シート生成 ------------------------------------------------
_SUMMARY_ORDER = [
    "入金", "給料賃金", "賞与", "法人税", "福利厚生費", "本部経費", "退職金積み立て",
    "旅費交通費", "広告宣伝費", "接待交際費", "賃借料", "通信費", "備品・消耗品費", "租税公課",
    "車両費", "水道光熱費", "保険料", "支払手数料", "研修採用費", "諸会費", "出金",
]
_SUMMARY_PL_NOTE = {"支払手数料": "→雑費", "研修採用費": "→その他経費", "諸会費": "→その他経費",
                    "法人税": "→PL対象外（市民税＝預り金）", "出金": "→PL対象外（小口補充・資金移動）"}


def keihi_workbook_bytes(bank: pd.DataFrame | None, petty: pd.DataFrame | None, sheet_name: str,
                         station: str = "桜木町", opening_petty: float | None = None) -> bytes:
    """経費Excel「YY.MM」シートと同じ配置（A〜G 銀行／I〜P 小口／R〜U 科目サマリ）を xlsx で返す。
    サマリは SUMIFS 数式で全行を参照する（手作業の集計範囲ズレを防ぐ）。科目は仕分け結果（店舗科目）を書く。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")
    ws["A1"] = f"{station}あおぞら銀行（法人口座）"
    ws["I1"] = "小口"
    hdr_bank = ["日付", "科目", "支払先", "摘要", "入金", "出金", "計"]
    hdr_petty = ["日付", "科目", "スタッフ", "支払先", "摘要", "入金", "出金", "計"]
    hdr_sum = ["科目", "銀行", "小口", "計", "PL行メモ"]
    for j, h in enumerate(hdr_bank, start=1):
        ws.cell(row=2, column=j, value=h).font = bold
    for j, h in enumerate(hdr_petty, start=9):
        ws.cell(row=2, column=j, value=h).font = bold
    for j, h in enumerate(hdr_sum, start=18):
        ws.cell(row=2, column=j, value=h).font = bold
    n_bank = 0
    if bank is not None and not bank.empty:
        r = 4
        for _, x in bank.iterrows():
            ws.cell(row=r, column=1, value=x.get("日付", ""))
            ws.cell(row=r, column=2, value=x.get("店舗科目", "") or x.get("科目", ""))
            ws.cell(row=r, column=3, value=x.get("支払先", "") or x.get("摘要", ""))
            ws.cell(row=r, column=4, value=x.get("摘要", "") if x.get("支払先", "") else "")
            ws.cell(row=r, column=5, value=float(x.get("入金", 0)) or None)
            ws.cell(row=r, column=6, value=float(x.get("出金", 0)) or None)
            ws.cell(row=r, column=7, value=f"=G{r-1}+E{r}-F{r}" if r > 4 else f"=E{r}-F{r}")
            if x.get("判定") in ("要確認", "判断不能", "要按分"):
                for c in range(1, 8):
                    ws.cell(row=r, column=c).fill = fill
            r += 1
        n_bank = r - 4
    n_petty = 0
    if petty is not None and not petty.empty:
        r = 4
        ws.cell(row=3, column=16, value=opening_petty if opening_petty is not None else None)
        for _, x in petty.iterrows():
            ws.cell(row=r, column=9, value=x.get("日付", ""))
            ws.cell(row=r, column=10, value=x.get("店舗科目", "") or x.get("科目", ""))
            ws.cell(row=r, column=11, value=x.get("スタッフ", ""))
            ws.cell(row=r, column=12, value=x.get("支払先", ""))
            ws.cell(row=r, column=13, value=x.get("摘要", ""))
            ws.cell(row=r, column=14, value=float(x.get("入金", 0)) or None)
            ws.cell(row=r, column=15, value=float(x.get("出金", 0)) or None)
            ws.cell(row=r, column=16, value=f"=P{r-1}+N{r}-O{r}")
            if x.get("判定") in ("要確認", "判断不能"):
                for c in range(9, 17):
                    ws.cell(row=r, column=c).fill = fill
            r += 1
        n_petty = r - 4
    last_b = max(4, 3 + n_bank)
    last_p = max(4, 3 + n_petty)
    r = 3
    for acc in _SUMMARY_ORDER:
        ws.cell(row=r, column=18, value=acc)
        if acc == "入金":
            fb = f"=SUMIFS($E$4:$E${last_b},$B$4:$B${last_b},R{r})"
            fp = f"=SUMIFS($N$4:$N${last_p},$J$4:$J${last_p},R{r})"
        else:
            fb = f"=SUMIFS($F$4:$F${last_b},$B$4:$B${last_b},R{r})-SUMIFS($E$4:$E${last_b},$B$4:$B${last_b},R{r})"
            fp = f"=SUMIFS($O$4:$O${last_p},$J$4:$J${last_p},R{r})-SUMIFS($N$4:$N${last_p},$J$4:$J${last_p},R{r})"
        ws.cell(row=r, column=19, value=fb)
        ws.cell(row=r, column=20, value=fp)
        ws.cell(row=r, column=21, value=f"=S{r}+T{r}")
        ws.cell(row=r, column=22, value=_SUMMARY_PL_NOTE.get(acc, ""))
        r += 1
    ws.cell(row=r, column=18, value="要按分/未確定")
    ws.cell(row=r, column=19, value=f'=SUMIFS($F$4:$F${last_b},$B$4:$B${last_b},"")+SUMIFS($F$4:$F${last_b},$B$4:$B${last_b},"要按分")')
    ws.cell(row=r, column=20, value=f'=SUMIFS($O$4:$O${last_p},$J$4:$J${last_p},"")')
    ws.cell(row=r, column=21, value=f"=S{r}+T{r}")
    r += 1
    ws.cell(row=r, column=18, value="計").font = bold
    ws.cell(row=r, column=19, value=f"=S3-SUM(S4:S{r-1})")
    ws.cell(row=r, column=20, value=f"=T3-SUM(T4:T{r-1})")
    ws.cell(row=r, column=21, value=f"=S{r}+T{r}")
    for col, w in {"A": 11, "B": 13, "C": 34, "D": 40, "E": 11, "F": 11, "G": 12, "I": 11, "J": 13, "K": 8,
                   "L": 26, "M": 34, "N": 9, "O": 9, "P": 10, "R": 14, "S": 12, "T": 11, "U": 12, "V": 22}.items():
        ws.column_dimensions[col].width = w
    for c in ("E", "F", "G", "N", "O", "P", "S", "T", "U"):
        for cell in ws[c]:
            cell.number_format = "#,##0"
    ws.freeze_panes = "A3"
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- 収支タブ貼り付けブロック ------------------------------------------------
def _memo_line(x: pd.Series) -> str:
    """収支タブのセルメモと同じタブ区切り（日付 科目 [スタッフ] 支払先 摘要 金額）。"""
    d = x.get("日付", "")
    d = d.strftime("%Y/%m/%d") if hasattr(d, "strftime") else str(d or "")[:10].replace("-", "/")
    parts = [d, str(x.get("店舗科目", "") or "")]
    if str(x.get("スタッフ", "") or ""):
        parts.append(str(x["スタッフ"]))
    parts.append(str(x.get("支払先", "") or ""))
    parts.append(str(x.get("摘要", "") or ""))
    amt = float(x.get("入金", 0)) if float(x.get("入金", 0)) else float(x.get("出金", 0))
    parts.append(f"{amt:,.0f}")
    return "\t".join(parts)


def pl_paste_block(bank: pd.DataFrame | None, petty: pd.DataFrame | None) -> pd.DataFrame:
    """PL行ごとに 金額 と セルメモ文（明細一覧）を返す。「 桜：収支」の実績列に値を、セルメモに明細を貼る。"""
    summ = pl_summary(bank, petty)
    parts = [d for d in (bank, petty) if d is not None and not d.empty]
    all_ = pd.concat(parts, ignore_index=True) if parts else pd.DataFrame()
    memos = []
    for pl in summ["PL行"]:
        if all_.empty:
            memos.append("")
            continue
        sub = all_[all_["PL行"] == pl]
        if pl == "入金":
            sub = sub[sub["入金"] > 0]
        memos.append("\n".join(_memo_line(x) for _, x in sub.iterrows()))
    summ["セルメモ"] = memos
    return summ
